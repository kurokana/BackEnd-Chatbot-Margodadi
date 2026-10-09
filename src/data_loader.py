from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class SourceRecord:
    file_name: str
    file_type: str
    source_path: str
    text: str
    metadata: dict[str, Any]


def load_pdf(path: Path) -> list[SourceRecord]:
    reader = PdfReader(str(path))
    records: list[SourceRecord] = []

    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = " ".join(text.split())
        if not text:
            continue

        records.append(
            SourceRecord(
                file_name=path.name,
                file_type="pdf",
                source_path=str(path),
                text=text,
                metadata={"page": page_index},
            )
        )

    return records


def _stringify_excel_row(sheet_name: str, row_number: int, row: pd.Series) -> str:
    parts = []
    for column, value in row.items():
        if pd.isna(value):
            continue

        value_text = str(value).strip()
        column_text = str(column).strip()
        if not value_text or column_text.startswith("Unnamed"):
            continue

        parts.append(f"{column_text}: {value_text}")

    return f"Data sheet {sheet_name}, baris {row_number}. " + ". ".join(parts)


def _normalize_excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.dropna(how="all")
    if frame.empty:
        return frame

    header_index = None
    for index, row in frame.iterrows():
        values = {str(value).strip().lower() for value in row.dropna().tolist()}
        if "no" in values and "nama usaha" in values:
            header_index = index
            break

    if header_index is None:
        return frame

    columns = frame.loc[header_index].fillna("").astype(str).str.strip().tolist()
    normalized = frame.loc[header_index + 1 :].copy()
    normalized.columns = columns
    return normalized.dropna(how="all")


def load_excel(path: Path) -> list[SourceRecord]:
    workbook = pd.read_excel(path, sheet_name=None, dtype=str, header=None)
    records: list[SourceRecord] = []

    for sheet_name, frame in workbook.items():
        frame = _normalize_excel_frame(frame)
        if frame.empty:
            continue

        for row_index, row in frame.iterrows():
            text = _stringify_excel_row(sheet_name, int(row_index) + 1, row)
            if text.strip().endswith("."):
                text = text.strip()
            else:
                text = text.strip() + "."

            if len(text) < 30:
                continue

            records.append(
                SourceRecord(
                    file_name=path.name,
                    file_type="xlsx",
                    source_path=str(path),
                    text=text,
                    metadata={"sheet": sheet_name, "row": int(row_index) + 1},
                )
            )

    return records


def load_docx(path: Path) -> list[SourceRecord]:
    document = Document(str(path))
    paragraphs = [
        " ".join(paragraph.text.split())
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.text.strip()
    ]

    table_texts: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
            if cells:
                table_texts.append(
                    f"Tabel {table_index}, baris {row_index}: " + " | ".join(cells)
                )

    text_blocks = paragraphs + table_texts
    records: list[SourceRecord] = []
    for block_index, text in enumerate(text_blocks, start=1):
        if len(text) < 20:
            continue

        records.append(
            SourceRecord(
                file_name=path.name,
                file_type="docx",
                source_path=str(path),
                text=text,
                metadata={"block": block_index},
            )
        )

    return records


def load_text(path: Path) -> list[SourceRecord]:
    text = " ".join(path.read_text(encoding="utf-8").split())
    if not text:
        return []

    return [
        SourceRecord(
            file_name=path.name,
            file_type="txt",
            source_path=str(path),
            text=text,
            metadata={"block": 1},
        )
    ]


def load_file(path: Path) -> list[SourceRecord]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix in {".xlsx", ".xls"}:
        return load_excel(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix == ".txt":
        return load_text(path)

    raise ValueError(f"Format file belum didukung: {suffix}")


def supported_file_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix in {".xlsx", ".xls"}:
        return suffix.lstrip(".")
    if suffix == ".docx":
        return "docx"
    if suffix == ".txt":
        return "txt"

    raise ValueError(f"Format file belum didukung: {suffix}")


def load_data_dir(data_dir: Path) -> list[SourceRecord]:
    records: list[SourceRecord] = []

    for path in sorted(data_dir.iterdir()):
        if not path.is_file():
            continue

        try:
            records.extend(load_file(path))
        except ValueError:
            continue

    return records
